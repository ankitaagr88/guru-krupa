"""Day book — the clinic's daily cash sheet — and the cash drawer (lane M).

The paper sheet had one row per patient: SR NO, DATE, NAME, PH NO, AGE ("20/F"), ADD (area), a money
column per kind of charge (OPD | MED | TEST | GLASSES | OT), TOTAL, MODE and LEFT (still owed); under
it the column totals, the opening "CASH BALANCE", cash taken out ("-3000 MAAM") and the closing
"CASH BAL". `day_book` builds the same thing from the bills:

  - one "visit" row per visit that day, ₹0 visits included (they show 0). Each bill line goes under
    its day-book column (`billing.line_head`); TOTAL is the bill total, MODE the modes of the money
    received that day, LEFT the balance still owed now;
  - one "old_balance" row per earlier visit's bill that received money that day (TOTAL = the money
    received; nothing under the columns — those charges were counted on the visit's own day);
  - one "ot" row per surgery that day (not cancelled) with a lens tier, OT team fees or a payment
    mode: the case total (lens tier price + the OT team's fees, app.services.ot_team) under the OT
    column (settings `otHead`). An OT case records only its lens tier, team fees and payment mode —
    no payment time or part payments — so the total counts on the surgery's date, as received when a
    payment mode is set and as LEFT otherwise. The lens price is the lens tier's current price (the
    case does not keep its own copy); team fees are the ones saved on the case.

The columns are the admin's day-book columns (`AccountHead`, Admin › Day book columns) in order;
a switched-off column only shows on a day that has money under it. A line whose column no longer
exists is counted under the settings' "other" column.

Cash drawer, per clinic day: closing = opening + cash received + cash put in − cash taken out. Only
cash-mode payments (bills and OT) touch the drawer; UPI / card / mediclaim are totals only. The
opening cash is the one typed for that day (`CashDay`) or else the closing cash carried forward
from the last day that had one typed (so Monday opens with Saturday's closing); a clinic that never
typed one starts at 0.
"""
import io
import re
from collections import defaultdict
from datetime import date, timedelta

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.db import utcnow
from app.models.audit import AuditLog
from app.models.billing import PAYMENT_MODES, AccountHead, BillItem, CashDay, CashMovement, StandardCharge
from app.models.config import ClinicSetting, LensTier
from app.models.ot import OtCase
from app.models.patients import Patient, Visit
from app.models.staff import Staff
from app.schemas.daybook import (AccountHeadOut, CashBox, CashMovementOut, DayBookOut, DayBookRow, DayBookSettings,
                                 DayBookTotals, HeadColumn, ModeAmount)
from app.services import billing, ot_team
from app.services.billing import BadValue, BillingError, Conflict, NotFound  # noqa: F401  (routes use them)
from app.services.reports import _aware, day_bounds, payments_on

SETTINGS_FIELDS = {"medicine_head": "medicines", "other_head": "hand-typed lines", "ot_head": "OT payments"}


def _audit(db: Session, by, action: str, entity: str, entity_id: int | None, **detail) -> None:
    db.add(AuditLog(staff_id=by.id if isinstance(by, Staff) else by, action=action, entity=entity,
                    entity_id=entity_id, detail=detail))


# --------------------------------------------------------------------------- columns (account heads)
def heads(db: Session, include_inactive: bool = False) -> list[AccountHead]:
    stmt = select(AccountHead).order_by(AccountHead.sort_order, AccountHead.id)
    if not include_inactive:
        stmt = stmt.where(AccountHead.active.is_(True))
    return list(db.scalars(stmt))


def head_use_counts(db: Session) -> dict[str, int]:
    counts: dict[str, int] = defaultdict(int)
    for model in (StandardCharge, BillItem):
        for key, n in db.execute(select(model.account_head_key, func.count()).where(model.account_head_key.is_not(None))
                                 .group_by(model.account_head_key)):
            counts[key] += n
    return counts


def head_out(head: AccountHead, counts: dict[str, int]) -> AccountHeadOut:
    return AccountHeadOut(id=head.id, key=head.key, label=head.label, sort_order=head.sort_order, active=head.active,
                          use_count=counts.get(head.key, 0))


def get_head(db: Session, key: str) -> AccountHead:
    head = db.scalar(select(AccountHead).filter_by(key=key))
    if head is None:
        raise NotFound("Day-book column not found")
    return head


def _label(db: Session, label: str, except_id: int | None = None) -> str:
    label = " ".join((label or "").split())
    if not label:
        raise BadValue("A column needs a name")
    q = select(AccountHead).where(func.lower(AccountHead.label) == label.lower())
    if except_id is not None:
        q = q.where(AccountHead.id != except_id)
    if db.scalar(q) is not None:
        raise Conflict(f"'{label}' is already a column")
    return label


def _key_from_label(label: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_")[:30] or "column"


def create_head(db: Session, *, label: str, key: str | None, by) -> AccountHead:
    label = _label(db, label)
    if key is None:
        base = key = _key_from_label(label)
        n = 2
        while db.scalar(select(AccountHead).filter_by(key=key)):
            key = f"{base[:27]}_{n}"
            n += 1
    elif db.scalar(select(AccountHead).filter_by(key=key)):
        raise Conflict(f"Column '{key}' already exists")
    head = AccountHead(key=key, label=label, active=True,
                       sort_order=(db.scalar(select(func.max(AccountHead.sort_order))) or 0) + 1)
    db.add(head)
    db.flush()
    _audit(db, by, "account_head.create", "account_head", head.id, key=key, label=label)
    db.commit()
    return head


def _settings_using(db: Session, key: str) -> list[str]:
    cfg = billing.daybook_settings(db)
    return [what for field, what in SETTINGS_FIELDS.items() if getattr(cfg, field) == key]


def update_head(db: Session, head: AccountHead, values: dict, by) -> AccountHead:
    changed = {}
    if values.get("label") is not None:
        label = _label(db, values["label"], head.id)
        if label != head.label:
            head.label = changed["label"] = label
    if values.get("active") is not None and values["active"] != head.active:
        if not values["active"] and (uses := _settings_using(db, head.key)):
            raise Conflict(f"{head.label} is the column for {' and '.join(uses)} — pick another one in the "
                           "settings below first")
        head.active = changed["active"] = values["active"]
    if changed:
        _audit(db, by, "account_head.update", "account_head", head.id, key=head.key, **changed)
        db.commit()
    return head


def delete_head(db: Session, head: AccountHead, by) -> None:
    used = head_use_counts(db).get(head.key, 0)
    if used:
        raise Conflict(f"{used} charge(s) or bill line(s) are counted under '{head.label}': switch it off instead")
    if uses := _settings_using(db, head.key):
        raise Conflict(f"{head.label} is the column for {' and '.join(uses)} — pick another one in the settings first")
    _audit(db, by, "account_head.delete", "account_head", head.id, key=head.key)
    db.delete(head)
    db.commit()


def reorder_heads(db: Session, keys: list[str], by) -> list[AccountHead]:
    rows = {h.key: h for h in db.scalars(select(AccountHead))}
    if set(keys) != set(rows) or len(keys) != len(rows):
        raise BadValue("order must list every existing column exactly once")
    for i, key in enumerate(keys):
        rows[key].sort_order = i
    _audit(db, by, "account_head.reorder", "account_head", None, order=keys)
    db.commit()
    return heads(db, include_inactive=True)


def save_settings(db: Session, values: dict, by) -> DayBookSettings:
    cfg = billing.daybook_settings(db)
    merged = cfg.model_dump()
    for field in SETTINGS_FIELDS:
        if values.get(field) is not None:
            head = db.scalar(select(AccountHead).filter_by(key=values[field]))
            if head is None:
                raise BadValue(f"Unknown day-book column '{values[field]}'")
            if not head.active:
                raise BadValue(f"'{head.label}' is switched off")
            merged[field] = values[field]
    new = DayBookSettings.model_validate(merged)
    row = db.get(ClinicSetting, billing.DAYBOOK_SETTINGS_KEY)
    value = new.model_dump(by_alias=True)
    if row is None:
        db.add(ClinicSetting(key=billing.DAYBOOK_SETTINGS_KEY, value=value))
    else:
        row.value = value
    _audit(db, by, "daybook_settings.update", "clinic_setting", None,
           **{k: v for k, v in values.items() if v is not None})
    db.commit()
    return new


# --------------------------------------------------------------------------- OT (lens + team) money
def _lens_prices(db: Session) -> dict[str, int]:
    return {t.key: t.price for t in db.scalars(select(LensTier))}


def _team_fees(case: OtCase) -> int:
    return ot_team.team_fees(ot_team.team_of(None, case))


def _ot_cases(db: Session, start: date, end: date) -> list[OtCase]:
    """Cases in [start, end] that carry money: a lens tier, OT team fees or a payment mode, not cancelled."""
    rows = db.scalars(select(OtCase).where(OtCase.date >= start, OtCase.date <= end, OtCase.status != "cancelled")
                      .order_by(OtCase.date, OtCase.id))
    return [c for c in rows
            if (c.billing or {}).get("lensTier") or (c.billing or {}).get("paymentMode") or _team_fees(c)]


def _ot_money(case: OtCase, prices: dict[str, int]) -> tuple[int, str | None]:
    """(case total = lens price + team fees, payment mode or None)."""
    b = case.billing or {}
    mode = b.get("paymentMode") or None
    total = prices.get(b.get("lensTier") or "", 0) + _team_fees(case)
    return total, (mode if mode in PAYMENT_MODES else None)


# --------------------------------------------------------------------------- cash drawer
def _cash_flow(db: Session, start: date, end: date) -> int:
    """Cash into the drawer over the clinic days [start, end): cash payments + OT cash + cash put in
    − cash taken out."""
    if end <= start:
        return 0
    bounds = (day_bounds(start)[0], day_bounds(end)[0])
    cash = sum(p.amount for p in payments_on(db, bounds) if p.mode == "cash")
    prices = _lens_prices(db)
    for case in _ot_cases(db, start, end - timedelta(days=1)):
        price, mode = _ot_money(case, prices)
        if mode == "cash":
            cash += price
    for m in db.scalars(select(CashMovement).where(CashMovement.day >= start, CashMovement.day < end)):
        cash += m.amount if m.direction == "in" else -m.amount
    return cash


def opening_for(db: Session, day: date) -> tuple[int, CashDay | None, str]:
    """(opening cash, the CashDay row if typed for this day, source: set | carried | none)."""
    row = db.get(CashDay, day)
    if row is not None:
        return row.opening_cash, row, "set"
    anchor = db.scalar(select(CashDay).where(CashDay.day < day).order_by(CashDay.day.desc()).limit(1))
    if anchor is None:
        return 0, None, "none"
    return anchor.opening_cash + _cash_flow(db, anchor.day, day), None, "carried"


def set_opening(db: Session, day: date, amount: int, note: str, by) -> None:
    row = db.get(CashDay, day)
    old = row.opening_cash if row is not None else opening_for(db, day)[0]
    if row is None:
        row = CashDay(day=day)
        db.add(row)
    row.opening_cash = amount
    row.note = (note or "").strip()
    row.set_by_staff_id = by.id if isinstance(by, Staff) else by
    row.set_at = utcnow()
    _audit(db, by, "cash_day.opening", "cash_day", None, day=day.isoformat(), old=old, new=amount, note=row.note)
    db.commit()


def add_movement(db: Session, day: date, direction: str, amount: int, person: str, reason: str, by) -> CashMovement:
    row = CashMovement(day=day, direction=direction, amount=amount, person=" ".join((person or "").split()),
                       reason=(reason or "").strip(), at=utcnow(), by_staff_id=by.id if isinstance(by, Staff) else by)
    db.add(row)
    db.flush()
    _audit(db, by, "cash_movement.create", "cash_movement", row.id, day=day.isoformat(), direction=direction,
           amount=amount, person=row.person, reason=row.reason)
    db.commit()
    return row


def remove_movement(db: Session, day: date, movement_id: int, by) -> None:
    row = db.get(CashMovement, movement_id)
    if row is None or row.day != day:
        raise NotFound("Cash entry not found")
    _audit(db, by, "cash_movement.delete", "cash_movement", row.id, day=day.isoformat(), direction=row.direction,
           amount=row.amount, person=row.person, reason=row.reason)
    db.delete(row)
    db.commit()


def people(db: Session, limit: int = 20) -> list[str]:
    """Names typed on earlier cash entries, most recent first (suggestions for "Who")."""
    seen: list[str] = []
    for (person,) in db.execute(select(CashMovement.person).where(CashMovement.person != "")
                                .order_by(CashMovement.at.desc(), CashMovement.id.desc())):
        if person.lower() not in (p.lower() for p in seen):
            seen.append(person)
        if len(seen) >= limit:
            break
    return seen


def _names(db: Session, ids) -> dict[int, str]:
    ids = {i for i in ids if i is not None}
    if not ids:
        return {}
    return {i: n for i, n in db.execute(select(Staff.id, Staff.name).where(Staff.id.in_(ids)))}


# --------------------------------------------------------------------------- the day's sheet
def _age_sex(patient: Patient | None, age=None, sex=None) -> str:
    if patient is not None:
        age, sex = patient.age, patient.sex
    parts = [str(age) if age is not None else "", sex or ""]
    return "/".join(parts) if any(parts) else ""


def day_book(db: Session, day: date) -> DayBookOut:
    cfg = billing.daybook_settings(db)
    all_heads = heads(db, include_inactive=True)
    known = {h.key for h in all_heads}
    bounds = day_bounds(day)
    payments = payments_on(db, bounds)
    paid_today: dict[int, list] = defaultdict(list)  # bill id -> payments that day
    for p in payments:
        paid_today[p.bill_id].append(p)

    def col(key: str) -> str:
        return key if key in known else cfg.other_head

    rows: list[DayBookRow] = []
    by_mode = {m: [0, 0] for m in PAYMENT_MODES}
    for p in payments:
        by_mode.setdefault(p.mode, [0, 0])
        by_mode[p.mode][0] += 1
        by_mode[p.mode][1] += p.amount

    # Visits that day, in registration order.
    visit_bills = set()
    for v in db.scalars(select(Visit).where(Visit.date == day).order_by(Visit.id)):
        bill = v.bill
        amounts: dict[str, int] = defaultdict(int)
        total = received = left = 0
        modes: list[str] = []
        status = ""
        if bill is not None:
            visit_bills.add(bill.id)
            for item in bill.items:
                amounts[col(billing.line_head(db, item, cfg))] += item.amount
            total, left, status = billing.bill_total(bill), billing.bill_balance(bill), billing.bill_status(bill)
            for p in paid_today.get(bill.id, []):
                received += p.amount
                if p.mode not in modes:
                    modes.append(p.mode)
        rows.append(DayBookRow(kind="visit", visit_id=v.id, patient_id=v.patient_id, name=v.patient.name,
                               phone=v.patient.phone, age_sex=_age_sex(v.patient), area=v.patient.address or "",
                               token=v.token, visit_date=v.date, amounts=dict(amounts), total=total,
                               received=received, modes=modes, left=left, status=status))

    # Balances from earlier visits collected that day.
    for bill_id, pays in paid_today.items():
        if bill_id in visit_bills:
            continue
        bill = pays[0].bill
        v = bill.visit
        modes = []
        for p in pays:
            if p.mode not in modes:
                modes.append(p.mode)
        received = sum(p.amount for p in pays)
        rows.append(DayBookRow(kind="old_balance", visit_id=v.id, patient_id=v.patient_id, name=v.patient.name,
                               phone=v.patient.phone, age_sex=_age_sex(v.patient), area=v.patient.address or "",
                               token=v.token, visit_date=v.date, amounts={}, total=received, received=received,
                               modes=modes, left=billing.bill_balance(bill), status=billing.bill_status(bill),
                               note=f"Old balance · visit of {v.date.strftime('%d %b %Y')}"))

    # Surgeries that day (lens price + OT team fees under the OT column).
    prices = _lens_prices(db)
    ot_col = col(cfg.ot_head)
    for case in _ot_cases(db, day, day):
        price, mode = _ot_money(case, prices)
        patient = case.patient
        received = price if mode else 0
        if mode:
            by_mode.setdefault(mode, [0, 0])
            by_mode[mode][0] += 1
            by_mode[mode][1] += price
        lens = (case.billing or {}).get("lensTier") or ""
        fees = _team_fees(case)
        rows.append(DayBookRow(kind="ot", ot_case_id=case.id, patient_id=case.patient_id,
                               name=patient.name if patient else case.patient_name,
                               phone=patient.phone if patient else None,
                               age_sex=_age_sex(patient, case.age, case.sex),
                               area=(patient.address or "") if patient else "", visit_date=case.date,
                               amounts={ot_col: price} if price else {}, total=price, received=received,
                               modes=[mode] if mode else [], left=price - received,
                               status="paid" if mode else ("unpaid" if price else ""),
                               note=" · ".join(x for x in (case.procedure, lens and f"lens {lens}",
                                                           fees and f"OT team ₹{fees:,}") if x)))

    # Columns: the active ones in order, plus a switched-off one that has money that day.
    used = defaultdict(int)
    for r in rows:
        for k, a in r.amounts.items():
            used[k] += a
    columns = [HeadColumn(key=h.key, label=h.label) for h in all_heads if h.active or used.get(h.key)]
    totals = DayBookTotals(amounts={c.key: used.get(c.key, 0) for c in columns}, total=sum(r.total for r in rows),
                           received=sum(r.received for r in rows), left=sum(max(r.left, 0) for r in rows))

    return DayBookOut(date=day, heads=columns, rows=rows, totals=totals,
                      by_mode=[ModeAmount(mode=m, payments=n, amount=a) for m, (n, a) in by_mode.items()],
                      received_total=sum(a for _, a in by_mode.values()),
                      cash=cash_box(db, day, by_mode.get("cash", [0, 0])[1]))


def cash_box(db: Session, day: date, cash_received: int) -> CashBox:
    opening, row, source = opening_for(db, day)
    moves = list(db.scalars(select(CashMovement).where(CashMovement.day == day).order_by(CashMovement.at,
                                                                                         CashMovement.id)))
    names = _names(db, [m.by_staff_id for m in moves] + [row.set_by_staff_id if row else None])
    cash_in = sum(m.amount for m in moves if m.direction == "in")
    cash_out = sum(m.amount for m in moves if m.direction == "out")
    return CashBox(opening_cash=opening, opening_source=source,
                   opening_set_by=names.get(row.set_by_staff_id) if row else None,
                   opening_set_at=_aware(row.set_at) if row else None, opening_note=row.note if row else "",
                   cash_received=cash_received,
                   movements=[CashMovementOut(id=m.id, direction=m.direction, amount=m.amount, person=m.person,
                                              reason=m.reason, at=_aware(m.at), by_name=names.get(m.by_staff_id))
                              for m in moves],
                   cash_in=cash_in, cash_out=cash_out, closing_cash=opening + cash_received + cash_in - cash_out)


# --------------------------------------------------------------------------- Excel download
MODE_LABELS = {"cash": "CASH", "upi": "UPI", "card": "CARD", "mediclaim": "MEDICLAIM"}


def xlsx(db: Session, day: date) -> bytes:
    """The day book laid out like the clinic's sheet: title "PATIENT LIST", the date, the opening
    cash, a header row, one row per patient, the totals, then the cash lines and totals by mode."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    book = day_book(db, day)
    wb = Workbook()
    ws = wb.active
    ws.title = day.strftime("%d-%m-%Y")
    header = ["SR NO", "DATE", "NAME", "PH NO", "AGE", "ADD", *[h.label for h in book.heads], "TOTAL", "MODE",
              "LEFT"]
    width = len(header)
    bold = Font(bold=True)
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(["PATIENT LIST"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append(["DATE", day.strftime("%d-%m-%Y")])
    ws.append(["CASH BALANCE", book.cash.opening_cash])
    for cell in (ws["A2"], ws["A3"]):
        cell.font = bold
    ws.append(header)
    head_row = ws.max_row
    for cell in ws[head_row]:
        cell.font = bold
        cell.border = box
        cell.alignment = Alignment(horizontal="center")

    date_txt = day.strftime("%d-%m-%Y")
    for i, r in enumerate(book.rows, start=1):
        name = r.name
        if r.kind == "old_balance":
            name = f"{r.name} (old balance {r.visit_date.strftime('%d-%m-%Y') if r.visit_date else ''})".strip()
        elif r.kind == "ot":
            name = f"{r.name} (OT)"
        mode = " + ".join(MODE_LABELS.get(m, m.upper()) for m in r.modes)
        ws.append([i, date_txt, name, r.phone or "", r.age_sex, r.area,
                   *[r.amounts.get(h.key, 0) for h in book.heads], r.total, mode, r.left if r.left else 0])
        for cell in ws[ws.max_row]:
            cell.border = box

    ws.append(["", "", "TOTAL", "", "", "", *[book.totals.amounts.get(h.key, 0) for h in book.heads],
               book.totals.total, "", book.totals.left])
    for cell in ws[ws.max_row]:
        cell.font = bold
        cell.border = box

    ws.append([])
    cash = book.cash
    ws.append(["CASH BALANCE (opening)", "", "", cash.opening_cash])
    ws.append(["+ CASH RECEIVED", "", "", cash.cash_received])
    for m in cash.movements:
        sign = "+" if m.direction == "in" else "-"
        label = f"{sign}{m.amount} {m.person}".strip() + (f" ({m.reason})" if m.reason else "")
        ws.append([label, "", "", m.amount if m.direction == "in" else -m.amount])
    ws.append(["CASH BAL", "", "", cash.closing_cash])
    ws.cell(row=ws.max_row, column=1).font = bold
    ws.cell(row=ws.max_row, column=4).font = bold
    ws.append([])
    ws.append(["RECEIVED BY MODE"])
    ws.cell(row=ws.max_row, column=1).font = bold
    for m in book.by_mode:
        ws.append([MODE_LABELS.get(m.mode, m.mode.upper()), "", "", m.amount])
    ws.append(["TOTAL RECEIVED", "", "", book.received_total])
    ws.cell(row=ws.max_row, column=1).font = bold

    widths = [7, 12, 28, 14, 8, 18, *[10] * len(book.heads), 10, 12, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

