"""Day book (the daily cash sheet), its columns in Admin and the cash drawer (lane M).

A made-up pair of clinic days in March 2024 (nobody else in the suite uses 2024), built straight
from the models so times and amounts are exact. IST = UTC + 5:30."""
import io
from datetime import date, datetime, timezone

import pytest

DAY = date(2024, 3, 11)
NEXT = date(2024, 3, 12)


def _utc(day, h, m=0):
    return datetime(day.year, day.month, day.day, h, m, tzinfo=timezone.utc)


@pytest.fixture(scope="module", autouse=True)
def _seed():
    from app.db import SessionLocal
    from app.seed.reference import seed_reference

    with SessionLocal() as db:
        seed_reference(db)


@pytest.fixture
def reception_headers(client, db):
    from app.auth.service import ensure_user

    ensure_user(db, username="dayrecep", password="rec123", name="Day Reception", role="reception")
    r = client.post("/api/auth/login", json={"username": "dayrecep", "password": "rec123"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture(scope="module")
def sheet():
    """DAY: Asha (OPD 500 + medicine 90 + glasses 1200, paid 1500 cash + 200 UPI, 90 left), Bharat
    (free follow-up, no bill), an OT case (lens tier, paid cash). NEXT: Asha pays her 90 in cash."""
    from app.db import SessionLocal
    from app.models.billing import Bill, BillItem, BillPayment, StandardCharge
    from app.models.config import LensTier
    from app.models.ot import OtCase
    from app.models.patients import Patient, Visit

    with SessionLocal() as db:
        opd = db.query(StandardCharge).filter(StandardCharge.account_head_key == "opd").first()
        tier = db.query(LensTier).order_by(LensTier.sort_order).first()
        asha = Patient(name="Daybook Asha", phone="9876500001", age=20, sex="F", address="Vesu")
        bharat = Patient(name="Daybook Bharat", phone="9876500002", age=61, sex="M", address="Adajan")
        db.add_all([asha, bharat])
        db.flush()
        va = Visit(patient_id=asha.id, date=DAY, token="#801", stage_key="done", status="completed",
                   created_at=_utc(DAY, 4))
        vb = Visit(patient_id=bharat.id, date=DAY, token="#802", stage_key="done", status="completed",
                   created_at=_utc(DAY, 5))
        db.add_all([va, vb])
        db.flush()
        db.add(Bill(visit_id=va.id, receipt_no="GK-2024-80001", payment_mode="upi", items=[
            # an older line with no column stored: worked out from its charge
            BillItem(label=opd.label, amount=500, kind="charge", standard_charge_id=opd.id),
            BillItem(label="Drops × 1", amount=90, kind="medicine"),  # -> settings medicineHead (med)
            BillItem(label="Glasses", amount=1200, kind="charge", account_head_key="glasses"),
        ], payments=[BillPayment(amount=1500, mode="cash", at=_utc(DAY, 6)),
                     BillPayment(amount=200, mode="upi", at=_utc(DAY, 6, 5)),
                     BillPayment(amount=90, mode="cash", at=_utc(NEXT, 5))]))
        case = OtCase(patient_id=asha.id, patient_name=asha.name, date=DAY, procedure="Cataract (phaco)",
                      status="completed", billing={"lensTier": tier.key, "mediclaim": False, "paymentMode": "cash"})
        db.add(case)
        db.commit()
        return {"asha": va.id, "bharat": vb.id, "lens": tier.price, "patient": asha.id}


def _book(client, headers, day):
    r = client.get(f"/api/daybook/{day.isoformat()}", headers=headers)
    assert r.status_code == 200, r.text
    return r.json()


def test_day_book_rows_columns_and_totals(client, admin_headers, sheet):
    book = _book(client, admin_headers, DAY)
    keys = [h["key"] for h in book["heads"]]
    assert keys[:6] == ["opd", "med", "test", "glasses", "ot", "other"]
    assert [h["label"] for h in book["heads"]][:5] == ["OPD", "MED", "TEST", "GLASSES", "OT"]

    visit_rows = {r["visitId"]: r for r in book["rows"] if r["kind"] == "visit"}
    a = visit_rows[sheet["asha"]]
    assert (a["name"], a["phone"], a["ageSex"], a["area"]) == ("Daybook Asha", "9876500001", "20/F", "Vesu")
    assert a["amounts"] == {"opd": 500, "med": 90, "glasses": 1200}
    assert (a["total"], a["received"], a["modes"]) == (1790, 1700, ["cash", "upi"])
    assert a["left"] == 0 and a["status"] == "paid"  # the 90 came in the next day (balance is as of now)
    b = visit_rows[sheet["bharat"]]  # a free visit with no bill still has its row, all 0
    assert (b["amounts"], b["total"], b["left"], b["modes"]) == ({}, 0, 0, [])

    ot = [r for r in book["rows"] if r["kind"] == "ot"]
    assert len(ot) == 1 and ot[0]["amounts"] == {"ot": sheet["lens"]} and ot[0]["modes"] == ["cash"]
    assert ot[0]["left"] == 0 and "Cataract" in ot[0]["note"]

    t = book["totals"]
    assert t["amounts"]["opd"] == 500 and t["amounts"]["glasses"] == 1200 and t["amounts"]["ot"] == sheet["lens"]
    assert t["total"] == 1790 + sheet["lens"]
    modes = {m["mode"]: m["amount"] for m in book["byMode"]}
    assert modes["cash"] == 1500 + sheet["lens"] and modes["upi"] == 200 and modes["card"] == 0
    assert book["receivedTotal"] == 1700 + sheet["lens"]


def test_old_balance_collected_later_shows_on_that_day(client, admin_headers, sheet):
    book = _book(client, admin_headers, NEXT)
    old = [r for r in book["rows"] if r["kind"] == "old_balance"]
    assert len(old) == 1
    assert (old[0]["visitId"], old[0]["total"], old[0]["received"], old[0]["modes"], old[0]["amounts"]) == (
        sheet["asha"], 90, 90, ["cash"], {})
    assert old[0]["visitDate"] == DAY.isoformat() and "Old balance" in old[0]["note"]
    assert book["cash"]["cashReceived"] == 90


def test_cash_drawer_arithmetic_and_opening_carried_forward(client, admin_headers, reception_headers, sheet, db):
    from app.models.audit import AuditLog

    d = DAY.isoformat()
    cash = _book(client, reception_headers, DAY)["cash"]
    assert cash["openingSource"] in ("none", "carried")

    r = client.put(f"/api/daybook/{d}/opening", json={"openingCash": 4040, "note": "counted"},
                   headers=reception_headers)
    assert r.status_code == 200, r.text
    cash = r.json()["cash"]
    assert (cash["openingCash"], cash["openingSource"], cash["openingSetBy"], cash["openingNote"]) == (
        4040, "set", "Day Reception", "counted")
    received = 1500 + sheet["lens"]
    assert cash["cashReceived"] == received and cash["closingCash"] == 4040 + received

    r = client.post(f"/api/daybook/{d}/movements", json={"direction": "out", "amount": 3000, "person": "MAAM",
                                                         "reason": "taken home"}, headers=reception_headers)
    assert r.status_code == 201, r.text
    client.post(f"/api/daybook/{d}/movements", json={"direction": "in", "amount": 500, "person": "Ramesh",
                                                    "reason": "change"}, headers=reception_headers)
    cash = _book(client, reception_headers, DAY)["cash"]
    assert (cash["cashOut"], cash["cashIn"]) == (3000, 500)
    assert cash["closingCash"] == 4040 + received + 500 - 3000
    assert [(m["direction"], m["person"], m["byName"]) for m in cash["movements"]] == [
        ("out", "MAAM", "Day Reception"), ("in", "Ramesh", "Day Reception")]
    assert client.post(f"/api/daybook/{d}/movements", json={"direction": "sideways", "amount": 1},
                       headers=reception_headers).status_code == 422
    assert client.post(f"/api/daybook/{d}/movements", json={"amount": 0}, headers=reception_headers).status_code == 422

    # the next day opens with this day's closing cash (not typed), plus that day's 90 cash
    closing = cash["closingCash"]
    nxt = _book(client, reception_headers, NEXT)["cash"]
    assert (nxt["openingCash"], nxt["openingSource"]) == (closing, "carried")
    assert nxt["closingCash"] == closing + 90
    # two days on (nothing happened on NEXT+1's eve): still carried
    later = _book(client, reception_headers, date(2024, 3, 14))["cash"]
    assert later["openingCash"] == closing + 90

    # undo the "in" entry; people typed before come back as suggestions
    in_id = cash["movements"][1]["id"]
    r = client.delete(f"/api/daybook/{d}/movements/{in_id}", headers=reception_headers)
    assert r.status_code == 200 and r.json()["cash"]["cashIn"] == 0
    assert client.delete(f"/api/daybook/{d}/movements/{in_id}", headers=reception_headers).status_code == 404
    assert "MAAM" in client.get("/api/daybook/people", headers=reception_headers).json()
    actions = {a.action for a in db.query(AuditLog).filter(AuditLog.entity.in_(["cash_day", "cash_movement"]))}
    assert {"cash_day.opening", "cash_movement.create", "cash_movement.delete"} <= actions


def test_xlsx_download_is_a_workbook_like_the_sheet(client, admin_headers, sheet):
    from openpyxl import load_workbook

    r = client.get(f"/api/daybook/{DAY.isoformat()}.xlsx", headers=admin_headers)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "day-book-2024-03-11.xlsx" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    cells = [tuple(row) for row in ws.iter_rows(values_only=True)]
    assert cells[0][0] == "PATIENT LIST" and cells[1][:2] == ("DATE", "11-03-2024")
    header = cells[3]
    assert header[:6] == ("SR NO", "DATE", "NAME", "PH NO", "AGE", "ADD") and header[-3:] == ("TOTAL", "MODE", "LEFT")
    assert "OPD" in header and "GLASSES" in header
    flat = [str(v) for row in cells for v in row if v is not None]
    assert "Daybook Asha" in flat and "CASH + UPI" in flat and "CASH BAL" in flat
    assert client.get("/api/daybook/not-a-date.xlsx", headers=admin_headers).status_code == 422


def test_day_book_roles(client, db):
    from app.auth.service import ensure_user

    ensure_user(db, username="dayot", password="ot123", name="Day OT", role="ot_staff")
    tok = client.post("/api/auth/login", json={"username": "dayot", "password": "ot123"}).json()["access_token"]
    h = {"Authorization": f"Bearer {tok}"}
    assert client.get(f"/api/daybook/{DAY.isoformat()}", headers=h).status_code == 403
    assert client.put(f"/api/daybook/{DAY.isoformat()}/opening", json={"openingCash": 1}, headers=h).status_code == 403
    assert client.get(f"/api/daybook/{DAY.isoformat()}").status_code == 401


# --------------------------------------------------------------------------- Admin › Day book columns
def test_head_admin_crud_reorder_and_delete_rules(client, admin_headers, reception_headers, sheet):
    base = "/api/admin/account-heads"
    assert client.post(base, json={"label": "Nope"}, headers=reception_headers).status_code == 403
    r = client.post(base, json={"label": "Contact Lens"}, headers=admin_headers)
    assert r.status_code == 201, r.text
    assert (r.json()["key"], r.json()["label"], r.json()["active"]) == ("contact_lens", "Contact Lens", True)
    assert client.post(base, json={"label": "contact lens"}, headers=admin_headers).status_code == 409

    r = client.patch(f"{base}/contact_lens", json={"label": "CL"}, headers=admin_headers)
    assert r.status_code == 200 and r.json()["label"] == "CL"
    client.patch(f"{base}/contact_lens", json={"active": False}, headers=admin_headers)
    assert "contact_lens" not in [h["key"] for h in client.get("/api/account-heads", headers=reception_headers).json()]
    everything = client.get("/api/account-heads", params={"includeInactive": "true"}, headers=admin_headers).json()
    assert "contact_lens" in [h["key"] for h in everything]

    # reorder: every key once
    keys = [h["key"] for h in everything]
    r = client.put(f"{base}/order", json={"keys": ["contact_lens"] + [k for k in keys if k != "contact_lens"]},
                   headers=admin_headers)
    assert r.status_code == 200 and r.json()[0]["key"] == "contact_lens"
    assert client.put(f"{base}/order", json={"keys": ["opd"]}, headers=admin_headers).status_code == 422
    client.put(f"{base}/order", json={"keys": keys}, headers=admin_headers)

    # in use (charges / bill lines) -> switch it off instead; in the settings -> can't switch off
    r = client.delete(f"{base}/glasses", headers=admin_headers)
    assert r.status_code == 409 and "switch it off" in r.json()["detail"]
    assert next(h for h in everything if h["key"] == "glasses")["useCount"] > 0
    r = client.patch(f"{base}/med", json={"active": False}, headers=admin_headers)
    assert r.status_code == 409 and "medicines" in r.json()["detail"]
    assert client.delete(f"{base}/contact_lens", headers=admin_headers).status_code == 204
    assert client.delete(f"{base}/contact_lens", headers=admin_headers).status_code == 404


def test_day_book_settings(client, admin_headers, reception_headers):
    cfg = client.get("/api/daybook-settings", headers=reception_headers).json()
    assert cfg == {"medicineHead": "med", "otherHead": "other", "otHead": "ot"}
    assert client.put("/api/admin/daybook-settings", json={"otherHead": "opd"},
                      headers=reception_headers).status_code == 403
    assert client.put("/api/admin/daybook-settings", json={"otherHead": "nope"},
                      headers=admin_headers).status_code == 422
    r = client.put("/api/admin/daybook-settings", json={"otherHead": "opd"}, headers=admin_headers)
    assert r.status_code == 200 and r.json()["otherHead"] == "opd" and r.json()["medicineHead"] == "med"
    # a typed line now defaults to OPD
    pid = client.post("/api/patients", json={"name": "Settings Patient", "age": 40, "sex": "M"},
                      headers=admin_headers).json()["id"]
    vid = client.post("/api/visits", json={"patientId": pid}, headers=admin_headers).json()["id"]
    b = client.put(f"/api/visits/{vid}/bill", json={"items": [{"label": "Typed", "amount": 10}]},
                   headers=reception_headers).json()
    assert b["items"][0]["accountHeadKey"] == "opd"
    client.put("/api/admin/daybook-settings", json={"otherHead": "other"}, headers=admin_headers)


def test_ot_row_counts_lens_plus_team_fees(client, admin_headers, db):
    """An OT case's day-book amount is its bill total: lens price + the OT team's fees. A case with only
    team fees (no lens, no payment mode yet) still shows, as money LEFT; cash counts in the drawer."""
    from app.models.config import LensTier
    from app.models.ot import OtCase

    day = date(2024, 3, 25)
    tier = db.query(LensTier).order_by(LensTier.sort_order).first()
    team = [{"roleKey": "surgeon", "roleLabel": "Surgeon", "name": "Dr. Own", "qualification": "MS", "regNo": "",
             "external": False, "partnerId": None, "fee": 0},
            {"roleKey": "anaesthetist", "roleLabel": "Anaesthetist", "name": "Dr. Visiting",
             "qualification": "MD Anaesthesia", "regNo": "", "external": True, "partnerId": None, "fee": 2500}]
    paid = OtCase(patient_name="Team Paid", date=day, procedure="Cataract (phaco)", status="completed",
                  billing={"lensTier": tier.key, "mediclaim": False, "paymentMode": "cash", "team": team})
    owed = OtCase(patient_name="Team Owed", date=day, procedure="Pterygium", status="scheduled",
                  billing={"lensTier": None, "mediclaim": False, "paymentMode": None,
                           "team": [dict(team[1], fee=1500)]})
    db.add_all([paid, owed])
    db.commit()

    book = _book(client, admin_headers, day)
    ot = {r["name"]: r for r in book["rows"] if r["kind"] == "ot"}
    p, o = ot["Team Paid"], ot["Team Owed"]
    assert (p["amounts"], p["total"], p["received"], p["left"]) == ({"ot": tier.price + 2500}, tier.price + 2500,
                                                                     tier.price + 2500, 0)
    assert "OT team ₹2,500" in p["note"]
    assert (o["amounts"], o["total"], o["received"], o["left"], o["status"]) == ({"ot": 1500}, 1500, 0, 1500, "unpaid")
    assert book["totals"]["amounts"]["ot"] == tier.price + 2500 + 1500
    assert book["cash"]["cashReceived"] == tier.price + 2500
